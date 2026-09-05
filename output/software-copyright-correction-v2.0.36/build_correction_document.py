from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SHOT_DIR = ROOT / "screenshots"
OUT = ROOT / "桌面歌词岛软件V2.0.36-文档鉴别材料补正稿.docx"

SOFTWARE = "桌面歌词岛软件"
VERSION = "V2.0.36"
OWNER = "么博丞"
BLUE = RGBColor(31, 78, 121)
INK = RGBColor(31, 41, 55)
MUTED = RGBColor(90, 100, 115)
LIGHT = "EAF1F8"
PALE = "F4F6F9"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[idx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_font(run, name="Microsoft YaHei", size=10.5, bold=None, color=INK):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    return run


def add_text(doc, text, *, size=10.5, bold=False, color=INK, align=None, before=0, after=6, line=1.25, keep=False):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    pf.keep_with_next = keep
    set_font(p.add_run(text), size=size, bold=bold, color=color)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    set_font(p.add_run(text), size={1: 16, 2: 13, 3: 11.5}[level], bold=True, color=BLUE)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.2
    set_font(p.add_run(text), size=10.2)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.2
    set_font(p.add_run(text), size=10.2)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(7)
    set_font(p.add_run(text), size=9, color=MUTED)
    return p


def add_figure(doc, filename, caption, width_cm=16.0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    run.add_picture(str(SHOT_DIR / filename), width=Cm(width_cm))
    add_caption(doc, caption)


def add_key_value_table(doc, rows, widths=(2500, 6860)):
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for label, value in rows:
        cells = table.add_row().cells
        set_cell_shading(cells[0], LIGHT)
        set_font(cells[0].paragraphs[0].add_run(label), size=9.6, bold=True, color=BLUE)
        set_font(cells[1].paragraphs[0].add_run(value), size=9.6)
        for cell in cells:
            cell.paragraphs[0].paragraph_format.space_after = Pt(2)
            cell.paragraphs[0].paragraph_format.line_spacing = 1.15
    set_table_geometry(table, list(widths))
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_page_header_footer(section):
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)
    set_font(hp.add_run(f"{SOFTWARE} {VERSION}｜文档鉴别材料补正"), size=8.5, color=MUTED)
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(fp.add_run(f"{SOFTWARE} {VERSION}  ·  第 "), size=8.5, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    fp._p.append(fld)
    set_font(fp.add_run(" 页"), size=8.5, color=MUTED)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for level, size, before, after in [(1, 16, 18, 10), (2, 13, 14, 7), (3, 11.5, 10, 5)]:
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLUE
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(10.2)
        style.paragraph_format.left_indent = Cm(0.95)
        style.paragraph_format.first_line_indent = Cm(-0.48)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.line_spacing = 1.2


def add_operation_page(doc, number, title, filename, caption, purpose, steps, evidence):
    add_heading(doc, f"{number}. {title}", 1)
    add_text(doc, purpose, size=10.3, after=5)
    add_figure(doc, filename, caption)
    add_heading(doc, "操作过程", 2)
    for item in steps:
        add_number(doc, item)
    add_heading(doc, "运行结果", 2)
    add_text(doc, evidence, size=10.2, after=0)
    doc.add_page_break()


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)
    configure_styles(doc)
    add_page_header_footer(section)

    add_text(doc, "计算机软件著作权登记", size=12, bold=True, color=BLUE, align=WD_ALIGN_PARAGRAPH.CENTER, before=40, after=12)
    add_text(doc, SOFTWARE, size=26, bold=True, color=INK, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
    add_text(doc, f"{VERSION} 文档鉴别材料补正稿", size=18, bold=True, color=BLUE, align=WD_ALIGN_PARAGRAPH.CENTER, after=24)
    add_text(doc, "软件运行功能、界面截图与运行数据补充说明", size=14, bold=True, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=28)
    add_key_value_table(doc, [
        ("软件全称", SOFTWARE),
        ("版本号", VERSION),
        ("软件简称", "歌词岛"),
        ("著作权人", OWNER),
        ("文档性质", "补正文档鉴别材料（实际运行界面与操作说明）"),
        ("截图日期", "2026年9月3日"),
    ])
    add_text(doc, "本补正稿针对审查意见补充软件全部主要功能的实际运行界面、操作步骤和运行数据。截图来自 V2.0.36 冻结版本的真实运行过程，保持原始界面比例，未进行界面重绘。", size=10.5, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, before=18, after=0)
    doc.add_page_break()

    add_heading(doc, "一、补正说明与版本对应", 1)
    add_text(doc, "本材料与原提交的软件说明书共同构成文档鉴别材料，用于补充说明桌面歌词岛软件 V2.0.36 的实际运行功能、操作路径和数据处理结果。各功能说明依照“启动播放器—读取系统媒体会话—检索与缓存歌词—同步渲染歌词岛—调整设置与布局”的连续使用流程编排。")
    add_key_value_table(doc, [
        ("冻结源代码提交", "439f35a7c2ffec426b7f11bf3fd57288ea4d1a0a"),
        ("内部构建标识", "2.0.36-Beta；对外申报版本统一为 V2.0.36"),
        ("源文件统计", "108 个自有 C# / XAML 文件"),
        ("源程序量", "13,108 个非空代码行，其中 C# 10,180 行、XAML 2,928 行"),
        ("鉴别材料截取", "前 1,500 行与后 1,500 行，共 3,000 行、60 页"),
        ("运行环境", "Windows 10.0.26200 x64；Microsoft Windows Desktop Runtime 3.1.32"),
        ("显示环境", "显示器 1，可用界面区域 1706 × 1018；顶部居中显示"),
    ])
    add_heading(doc, "二、完整功能操作链", 1)
    for item in [
        "启动软件，完成单实例检查、偏好设置加载、媒体会话服务和歌词缓存初始化。",
        "在 Apple Music 等兼容播放器中播放歌曲，软件读取曲名、歌手、封面、播放状态和时间轴。",
        "软件按首选歌词源检索候选，解析同步歌词与翻译，并把结果写入本地缓存。",
        "顶部歌词岛显示封面、当前歌词、下一句/翻译和播放控制，并随时间轴持续更新。",
        "进入偏好设置调整歌词、屏幕位置、自动收起、缓存、鼠标避让、快捷键与模块布局。",
        "点击“应用”即时生效；点击“保存”持久化设置；点击“取消”放弃未保存的修改。",
    ]:
        add_number(doc, item)
    add_heading(doc, "三、截图覆盖范围", 1)
    add_text(doc, "后续页面依次给出顶部歌词岛运行态及全部主要设置页面。每张截图后均说明操作和可观察结果，以证明界面、数据和功能之间的对应关系。", after=0)
    doc.add_page_break()

    add_operation_page(
        doc, 1, "顶部歌词岛、歌词同步与播放控制", "00-main-runtime.png",
        "图 1  V2.0.36 顶部歌词岛真实运行界面（封面、双行歌词、上一曲/暂停/下一曲）",
        "软件连接 Windows 系统媒体会话后，在屏幕顶部显示当前播放内容，并按照播放器时间轴更新同步歌词。",
        [
            "启动桌面歌词岛软件并保持后台运行。",
            "在兼容播放器中开始播放音乐，等待软件读取媒体元数据并完成歌词检索。",
            "观察顶部歌词岛：左侧显示专辑封面，中部显示当前句与下一句，右侧显示播放控制按钮。",
            "点击上一曲、播放/暂停或下一曲；软件只对播放器声明支持的控制能力发送命令。",
        ],
        "实际运行时识别到同步歌词“布拉格的广场拥挤的剧场 / 安静小巷一家咖啡馆”，封面和三个播放控制按钮同时显示，证明媒体会话读取、歌词解析、时间轴选择、模块渲染和播放控制链路已运行。",
    )

    add_operation_page(
        doc, 2, "歌词显示、歌词来源与主题", "01-lyrics-display.png",
        "图 2  歌词显示设置：歌词源、单行/多行、翻译与主题",
        "用户可以选择自动或指定歌词来源，切换单行/多行显示，控制歌词库已有中文翻译，并选择浅色、深色或跟随系统主题。",
        [
            "打开偏好设置，进入“歌词显示”。",
            "在“首选歌词源”中选择自动或指定来源；自动模式会在来源不可用时继续尝试其他来源。",
            "在“行数”中选择单行或多行；按需要勾选“显示歌词库里的中文翻译”。",
            "在左下角选择浅色、深色或跟随系统，点击“应用”预览，确认后点击“保存”。",
        ],
        "截图中首选歌词源为“自动选择”，已选择“多行”并开启中文翻译，主题为跟随系统。设置与图 1 的双行歌词显示相互对应。",
    )

    add_operation_page(
        doc, 3, "显示器定位、边缘位置与自动收起", "02-position-status.png",
        "图 3  位置与状态设置：显示器、边缘位置和计时参数",
        "软件支持多显示器选择、屏幕顶边水平定位，并可配置无播放后的收起时间和自动折叠布局的展开停留时间。",
        [
            "进入“位置与状态”，在显示器列表中选择目标屏幕。",
            "拖动“边缘位置”滑块，或点击“居中”，确定歌词岛在屏幕顶边的位置。",
            "设置“无播放后收起”秒数；停止播放后达到设定时间，歌词岛滑出主要可见区域。",
            "设置“展开停留”秒数；自动折叠布局在交互结束后按该时间恢复紧凑状态。",
        ],
        "实际显示器为“显示器 1（1706 × 1018）”，歌词岛位于顶部居中；无播放后 6 秒收起，展开停留时间为 5 秒。",
    )

    add_operation_page(
        doc, 4, "同步歌词缓存与容量控制", "03-cache.png",
        "图 4  缓存设置：容量上限、复用方式与淘汰策略",
        "软件把已下载的同步歌词保存在本地，下次播放同一曲目时优先读取缓存，以减少等待和重复网络请求。",
        [
            "进入“缓存”，输入缓存容量（MB）。",
            "播放一首未缓存歌曲，软件检索歌词并写入本地歌词文件。",
            "再次播放相同曲目，软件优先读取缓存；超过容量上限时按最久未使用顺序清理。",
            "点击“应用”或“保存”，容量限制立即参与后续写入和清理。",
        ],
        "截图中的缓存上限为 64 MB。本次实际运行共产生 6 个歌词缓存文件、合计 21,523 字节，说明网络检索结果已落盘并可被后续播放复用。",
    )

    add_operation_page(
        doc, 5, "鼠标避让、透明光晕与点击穿透", "04-mouse-avoidance.png",
        "图 5  鼠标避让设置与实时预览",
        "鼠标靠近歌词岛时，软件按探测范围和光晕频谱降低局部不透明度，使下方窗口内容保持可读；可选穿透允许左键操作下方窗口。",
        [
            "进入“鼠标避让”，调整光晕大小、探测范围和光晕形状。",
            "在实时预览区域移动鼠标，观察歌词岛局部变淡。",
            "调整过渡位置及中心、过渡、边缘透明度，形成连续透明梯度。",
            "按需要开启“透明避让时允许左键点到下方窗口”，并点击“保存”。",
        ],
        "实际参数为光晕 86 px、探测范围 60 px、形状 1.27:1、过渡位置 56%，中心/过渡/边缘透明度分别为 98%/97%/0%，并已启用点击穿透。",
    )

    add_operation_page(
        doc, 6, "全局快捷键与歌词时间偏移", "05-hotkeys.png",
        "图 6  快捷键设置：提前、延后、重置与临时交互",
        "软件提供全局快捷键校准歌词时间轴，并允许按住临时交互键展开折叠布局或暂停点击穿透。",
        [
            "进入“快捷键”，单击需要修改的输入框并按下新的组合键。",
            "播放期间按 Ctrl+Alt+Left，使歌词提前 0.5 秒；按 Ctrl+Alt+Right，使歌词延后 0.5 秒。",
            "按 Ctrl+Alt+Down，将歌词偏移恢复为默认值。",
            "按住 Ctrl 临时启用交互，完成歌词岛控制后松开。",
        ],
        "截图显示四项快捷键均已注册。运行时偏移会参与当前歌词行计算，用于播放器时间轴存在小误差时的人工校准。",
    )

    add_operation_page(
        doc, 7, "模块布局、播放器选择与真实布局编辑", "06-module-layout.png",
        "图 7  模块布局设置：布局模式、播放器、模块工具箱与样式参数",
        "V2.0.36 提供水平积木和自动折叠两种布局。用户可选择播放器，并在真实歌词岛中添加、排序或删除歌词、封面、播放、信息、进度和分割线模块。",
        [
            "进入“模块布局”，选择“水平积木”或“自动折叠”。",
            "在“播放器”中选择自动模式或锁定已检测到的播放器；自动模式跟随最近活跃会话。",
            "从模块工具箱把歌词、封面、播放、信息、进度或分割线拖到屏幕顶部真实歌词岛。",
            "在岛内拖动以排序，拖出岛外以删除；调整歌词宽度、分割线透明度和左右间距。",
            "点击“保存”提交布局；点击“取消”或关闭窗口时回滚未提交草稿。",
        ],
        "截图中水平积木布局被选中，播放器为自动选择，六类模块全部可用；歌词宽度为 520 px，分割线透明度 22%、左右间距 4 px。",
    )

    add_operation_page(
        doc, 8, "版本信息、教程与兼容范围", "07-about-version.png",
        "图 8  关于页面：V2.0.36 版本号及主要更新内容",
        "关于页面用于确认软件名称、版本、作者和本版本主要能力，并提供项目主页及重新进入教学模式的入口。",
        [
            "进入“关于”，核对软件名称和版本号。",
            "阅读 V2.0 Beta 更新内容，确认布局、模块、播放器和缓存等能力均属于本版本。",
            "需要重新熟悉操作时点击“重新开始教学”，按照遮罩提示完成播放、设置和布局操作。",
        ],
        "页面明确显示版本号“v2.0.36 Beta”，并列出水平积木/自动折叠、六类模块、多播放器 SMTC、同步歌词缓存、播放器锁定、鼠标避让、快捷键和主题设置。",
    )

    add_operation_page(
        doc, 9, "支持与反馈入口", "08-support-developer.png",
        "图 9  支持开发者页面：评价、分享、GitHub、反馈与 Pro 入口",
        "除核心歌词功能外，软件提供评价、分享、项目主页、意见反馈和支持计划入口；这些入口不影响主体歌词功能免费使用。",
        [
            "进入“支持开发者”，查看免费支持和 Pro 支持计划。",
            "用户可选择评价、分享或访问 GitHub；未上线的反馈入口以状态文字提示，不执行不可用操作。",
            "支持计划页面展示功能说明和购买入口；是否使用由用户自行决定。",
        ],
        "页面完整显示四类免费支持入口及 Pro 支持计划说明。该页面属于辅助功能，不参与歌词检索、同步显示和媒体控制的核心处理链路。",
    )

    add_heading(doc, "十、实际运行数据与功能对应", 1)
    add_text(doc, "以下数据来自本次 V2.0.36 实际运行过程，用于说明界面操作确实触发了媒体、歌词、缓存和设置处理。")
    add_key_value_table(doc, [
        ("媒体会话", "兼容播放器播放状态可被读取；歌词岛显示封面、歌词和控制按钮"),
        ("同步歌词", "运行界面显示双行同步歌词；设置中启用多行和歌词库中文翻译"),
        ("缓存文件数", "6 个"),
        ("缓存总大小", "21,523 字节"),
        ("缓存样例 1", "sabrina-carpenter-nonsense-163.lrc，6,151 字节"),
        ("缓存样例 2", "zara-larsson-lush-life-200.lrc，6,033 字节"),
        ("缓存样例 3", "jolin-294.lrc，4,026 字节"),
        ("显示与定位", "显示器 1（1706 × 1018），屏幕顶边居中"),
        ("缓存上限", "64 MB；超过上限时优先清理最久未使用文件"),
        ("歌词偏移", "默认 800 ms；支持每次 ±500 ms 调整及一键重置"),
        ("鼠标避让", "光晕 86 px，探测 60 px，中心透明度 98%"),
    ])
    add_heading(doc, "十一、主要功能与申请表对应关系", 1)
    mapping = [
        ("系统媒体会话与多播放器", "图 1、图 7、图 8", "读取播放状态、选择或锁定播放器、按能力控制"),
        ("歌词检索、翻译与同步", "图 1、图 2、运行数据", "多来源检索、同步行选择、中文翻译、时间偏移"),
        ("模块化顶部歌词岛", "图 1、图 7", "封面、歌词、播放、信息、进度、分割线组合"),
        ("真实布局编辑", "图 7", "拖入添加、岛内排序、拖出删除、保存或取消回滚"),
        ("位置、自动收起与多显示器", "图 3", "屏幕选择、顶边位置、暂停/无播放收起"),
        ("鼠标避让与点击穿透", "图 5", "透明光晕、探测范围、实时预览、穿透"),
        ("缓存和本地设置", "图 4、运行数据", "歌词缓存、容量限制、最近最少使用清理、设置持久化"),
        ("快捷键校准", "图 6", "提前、延后、重置偏移、临时交互"),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for idx, label in enumerate(("申请表主要功能", "截图/数据证据", "软件实际处理")):
        set_cell_shading(table.rows[0].cells[idx], LIGHT)
        set_font(table.rows[0].cells[idx].paragraphs[0].add_run(label), size=9, bold=True, color=BLUE)
    for row in mapping:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_font(cells[idx].paragraphs[0].add_run(value), size=8.8)
            cells[idx].paragraphs[0].paragraph_format.space_after = Pt(1)
            cells[idx].paragraphs[0].paragraph_format.line_spacing = 1.1
    set_table_geometry(table, [2850, 1850, 4660])
    add_text(doc, "结论：本补正材料所示界面、操作流程、运行数据及代码规模均对应桌面歌词岛软件 V2.0.36，能够连续、完整地反映申请表所述主要功能和技术特点。", size=10.5, bold=True, color=BLUE, before=12, after=0)

    core = doc.core_properties
    core.title = f"{SOFTWARE}{VERSION}-文档鉴别材料补正稿"
    core.subject = "软件运行功能、界面截图与运行数据补充说明"
    core.author = OWNER
    core.keywords = "计算机软件著作权, 补正, V2.0.36, 歌词岛"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
